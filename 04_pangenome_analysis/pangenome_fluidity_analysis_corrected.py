#!/usr/bin/env python3
"""
Publication-oriented pangenome/genomic fluidity analysis.

The script calculates the genomic fluidity statistic described by Kislyuk et al.
(2011) from a Panaroo/Roary binary gene_presence_absence.Rtab matrix.

Primary outputs
---------------
- Overall and group-specific fluidity estimates
- Genome-level bootstrap and delete-one-genome jackknife confidence intervals
- Pairwise within/between-group summaries
- Rare-gene sensitivity analyses
- Sample-size convergence analyses
- Optional two-group permutation test and equal-size resampling
- Publication-quality PNG, PDF and SVG figures
- A publication-ready Methods/Results text file
- A DOCX results report

The core-genome alignment is not required for the fluidity calculation. When an
alignment is supplied, the script audits sample identifiers against the Rtab
matrix but does not use SNP variation to calculate fluidity.
"""

from __future__ import annotations

import argparse
import logging
import math
import re
import sys
from collections import Counter
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Dict, Iterable, List, Optional, Sequence, Tuple

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from scipy import sparse

try:
    from docx import Document
    from docx.enum.section import WD_SECTION
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt
    DOCX_AVAILABLE = True
except ImportError:
    DOCX_AVAILABLE = False


LOGGER = logging.getLogger("pangenome_fluidity")


@dataclass
class FluidityResult:
    label: str
    n_genomes: int
    n_gene_clusters: int
    n_pairs: int
    mean_phi: float
    median_phi: float
    q1_phi: float
    q3_phi: float
    min_phi: float
    max_phi: float
    jackknife_se: float
    jackknife_ci_low: float
    jackknife_ci_high: float
    bootstrap_mean: float
    bootstrap_ci_low: float
    bootstrap_ci_high: float


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Calculate publication-oriented pangenome/genomic fluidity from a Panaroo Rtab matrix.",
        formatter_class=argparse.ArgumentDefaultsHelpFormatter,
    )
    parser.add_argument("--rtab", required=True, type=Path,
                        help="Panaroo/Roary gene_presence_absence.Rtab file.")
    parser.add_argument("--metadata", type=Path,
                        help="Optional tab-delimited metadata file.")
    parser.add_argument("--alignment", type=Path,
                        help="Optional core-genome FASTA alignment used only for identifier auditing.")
    parser.add_argument("--output", type=Path, default=Path("fluidity_results"),
                        help="Output directory.")
    parser.add_argument("--id-column",
                        help="Metadata column containing genome/sample identifiers. Auto-detected if omitted.")
    parser.add_argument("--group-column",
                        help="Metadata column defining analytical groups (for this dataset use isolate_class). Auto-detected only from sufficiently populated columns if omitted.")
    parser.add_argument("--default-group", default="All",
                        help="Group assigned when no usable metadata grouping column is available.")
    parser.add_argument("--compare-groups", nargs=2, metavar=("GROUP_A", "GROUP_B"),
                        help="Two metadata groups for permutation testing and equal-size resampling.")
    parser.add_argument("--organism", default="Pseudomonas aeruginosa",
                        help="Organism name used in the generated report.")
    parser.add_argument("--bootstrap", type=int, default=1000,
                        help="Genome-level bootstrap replicates per reported dataset; set 0 to disable.")
    parser.add_argument("--bootstrap-batch", type=int, default=100,
                        help="Number of bootstrap replicates processed per matrix batch.")
    parser.add_argument("--permutations", type=int, default=999,
                        help="Label permutations for a two-group comparison; set 0 to disable.")
    parser.add_argument("--matched-replicates", type=int, default=1000,
                        help="Equal-sample-size resampling replicates for the selected two groups.")
    parser.add_argument("--convergence-replicates", type=int, default=100,
                        help="Random subsets per sample size for convergence analysis.")
    parser.add_argument("--seed", type=int, default=20260806,
                        help="Random seed for reproducibility.")
    parser.add_argument("--save-pairwise", action="store_true",
                        help="Write all pairwise fluidity values to a compressed TSV file.")
    parser.add_argument("--max-plot-pairs", type=int, default=100000,
                        help="Maximum pairwise values sampled per category for violin plots.")
    parser.add_argument("--min-group-size", type=int, default=5,
                        help="Minimum group size for group-specific estimates.")
    return parser.parse_args()


def configure_logging(output_dir: Path) -> None:
    output_dir.mkdir(parents=True, exist_ok=True)
    LOGGER.setLevel(logging.INFO)
    formatter = logging.Formatter("%(asctime)s | %(levelname)s | %(message)s")
    stream = logging.StreamHandler(sys.stdout)
    stream.setFormatter(formatter)
    file_handler = logging.FileHandler(output_dir / "analysis.log", mode="w")
    file_handler.setFormatter(formatter)
    LOGGER.handlers.clear()
    LOGGER.addHandler(stream)
    LOGGER.addHandler(file_handler)


def normalize_id(value: object) -> str:
    """Normalize common file-derived sample identifiers without removing accession versions."""
    text = str(value).strip().strip('"').strip("'")
    text = Path(text).name
    suffixes = [".gz", ".bz2", ".xz", ".gff3", ".gff", ".fasta", ".fna", ".fa",
                ".ffn", ".faa", ".gbff", ".gbk", ".genbank", ".aln"]
    changed = True
    while changed:
        changed = False
        lower = text.lower()
        for suffix in suffixes:
            if lower.endswith(suffix):
                text = text[:-len(suffix)]
                changed = True
                break
    return text.strip()


def read_rtab(path: Path) -> Tuple[np.ndarray, List[str], List[str]]:
    if not path.exists():
        raise FileNotFoundError(f"Rtab file not found: {path}")
    LOGGER.info("Reading binary gene presence/absence matrix: %s", path)
    try:
        header = pd.read_csv(path, sep="\t", nrows=0).columns.tolist()
        if len(header) < 3:
            raise ValueError("Rtab header contains fewer than two genome columns.")
        dtype_map = {column: np.uint8 for column in header[1:]}
        frame = pd.read_csv(path, sep="\t", index_col=0, dtype=dtype_map, low_memory=False)
    except (TypeError, ValueError):
        LOGGER.warning("Direct uint8 import failed; coercing Rtab entries to numeric values.")
        frame = pd.read_csv(path, sep="\t", index_col=0, low_memory=False)
        frame = frame.apply(pd.to_numeric, errors="coerce").fillna(0)

    if frame.empty or frame.shape[1] < 2:
        raise ValueError("The Rtab matrix is empty or contains fewer than two genome columns.")
    if frame.columns.duplicated().any():
        duplicates = frame.columns[frame.columns.duplicated()].tolist()
        raise ValueError(f"Duplicate genome columns in Rtab: {duplicates[:10]}")

    genome_names = [str(c).strip() for c in frame.columns]
    gene_names = [str(i).strip() for i in frame.index]
    values = frame.to_numpy(copy=False)
    non_binary = np.logical_and(values != 0, values != 1)
    if np.any(non_binary):
        count = int(non_binary.sum())
        LOGGER.warning("Found %d non-binary entries; values greater than zero will be treated as presence.", count)
    matrix = (values > 0).astype(np.uint8, copy=False)

    zero_rows = np.where(matrix.sum(axis=1) == 0)[0]
    if zero_rows.size:
        LOGGER.warning("Found %d all-zero gene-cluster rows; they do not affect fluidity and will be retained in counts.",
                       zero_rows.size)
    zero_cols = np.where(matrix.sum(axis=0) == 0)[0]
    if zero_cols.size:
        names = [genome_names[i] for i in zero_cols[:10]]
        raise ValueError(f"Genome columns with zero detected genes were found: {names}")

    LOGGER.info("Loaded %d gene clusters across %d genomes.", matrix.shape[0], matrix.shape[1])
    return matrix, gene_names, genome_names


def infer_metadata_id_column(metadata: pd.DataFrame, genome_names: Sequence[str]) -> Tuple[Optional[str], int]:
    normalized_genomes = set(normalize_id(x) for x in genome_names)
    best_col: Optional[str] = None
    best_overlap = 0
    for col in metadata.columns:
        values = metadata[col].dropna().astype(str)
        if values.empty:
            continue
        normalized_values = set(normalize_id(x) for x in values)
        overlap = len(normalized_values & normalized_genomes)
        if overlap > best_overlap:
            best_col, best_overlap = str(col), overlap
    return best_col, best_overlap


def infer_group_column(metadata: pd.DataFrame, excluded: Optional[str] = None) -> Optional[str]:
    """Infer a biologically meaningful, sufficiently populated grouping column.

    Sparse metadata fields must not be selected merely because their names look
    suitable. A candidate is accepted only when at least 80% of rows contain a
    non-missing group value after canonicalisation. The curated ``isolate_class``
    field is prioritised because it records the analytical classification used
    to construct the present clinical collection.
    """
    preferred = [
        "isolate_class", "source_category", "classification", "source_type",
        "source", "origin", "group", "category", "isolation_category",
        "sample_source", "host_category", "environment"
    ]
    lower_map = {str(c).lower(): str(c) for c in metadata.columns}

    def usable(column: str) -> bool:
        values = metadata[column].map(canonicalize_group)
        nonmissing = values.ne("Unknown")
        coverage = float(nonmissing.mean()) if len(values) else 0.0
        unique_count = int(values[nonmissing].nunique())
        return coverage >= 0.80 and 1 <= unique_count <= 50

    for candidate in preferred:
        if candidate in lower_map:
            column = lower_map[candidate]
            if column != excluded and usable(column):
                return column

    keywords = ("clinical", "environment", "environmental", "other", "human", "patient")
    best_col: Optional[str] = None
    best_score = -1.0
    for col in metadata.columns:
        col = str(col)
        if col == excluded:
            continue
        canonical = metadata[col].map(canonicalize_group)
        nonmissing = canonical.ne("Unknown")
        coverage = float(nonmissing.mean()) if len(canonical) else 0.0
        unique_count = int(canonical[nonmissing].nunique())
        if coverage < 0.80 or unique_count < 1 or unique_count > 50:
            continue
        raw = metadata[col].fillna("").astype(str).str.lower()
        keyword_fraction = float(raw.apply(lambda x: any(k in x for k in keywords)).mean())
        score = coverage + keyword_fraction
        if score > best_score:
            best_col, best_score = col, score
    return best_col


def canonicalize_group(value: object) -> str:
    text = str(value).strip()
    low = text.lower()
    if not text or low in {"nan", "none", "na", "n/a", "unknown", "not available"}:
        return "Unknown"
    if "clinical" in low:
        return "Clinical"
    if "environment" in low:
        return "Environmental"
    if low == "other" or low.startswith("other "):
        return "Other"
    return text


def map_metadata_groups(
    metadata_path: Optional[Path],
    genome_names: Sequence[str],
    id_column: Optional[str],
    group_column: Optional[str],
    default_group: str,
    output_dir: Path,
) -> Tuple[np.ndarray, Optional[str], Optional[str], pd.DataFrame]:
    audit = pd.DataFrame({
        "rtab_genome": genome_names,
        "normalized_rtab_id": [normalize_id(x) for x in genome_names],
    })

    if metadata_path is None:
        LOGGER.info("No metadata file supplied; assigning all genomes to '%s'.", default_group)
        audit["metadata_match"] = False
        audit["group"] = default_group
        audit.to_csv(output_dir / "metadata_id_audit.tsv", sep="\t", index=False)
        return np.array([default_group] * len(genome_names), dtype=object), None, None, audit

    if not metadata_path.exists():
        raise FileNotFoundError(f"Metadata file not found: {metadata_path}")
    metadata = pd.read_csv(metadata_path, sep="\t", dtype=str, low_memory=False)
    if metadata.empty:
        raise ValueError("The metadata file is empty.")

    if id_column is not None and id_column not in metadata.columns:
        raise ValueError(f"Requested metadata ID column '{id_column}' was not found.")
    if id_column is None:
        id_column, overlap = infer_metadata_id_column(metadata, genome_names)
        if id_column is None or overlap == 0:
            LOGGER.warning("No metadata ID column overlapped the Rtab headers; grouping will use '%s'.", default_group)
            audit["metadata_match"] = False
            audit["group"] = default_group
            audit.to_csv(output_dir / "metadata_id_audit.tsv", sep="\t", index=False)
            return np.array([default_group] * len(genome_names), dtype=object), None, None, audit
        LOGGER.info("Auto-detected metadata ID column '%s' with %d/%d Rtab matches.",
                    id_column, overlap, len(genome_names))

    if group_column is not None and group_column not in metadata.columns:
        raise ValueError(f"Requested group column '{group_column}' was not found.")
    if group_column is None:
        group_column = infer_group_column(metadata, excluded=id_column)
        if group_column:
            LOGGER.info("Auto-detected metadata grouping column '%s'.", group_column)
        else:
            LOGGER.warning("No grouping column was detected; matched genomes will use '%s'.", default_group)

    metadata = metadata.copy()
    metadata["__normalized_id"] = metadata[id_column].map(normalize_id)
    duplicate_mask = metadata["__normalized_id"].duplicated(keep=False)
    if duplicate_mask.any():
        duplicates = metadata.loc[duplicate_mask, [id_column, "__normalized_id"]]
        duplicates.to_csv(output_dir / "metadata_duplicate_ids.tsv", sep="\t", index=False)
        LOGGER.warning("Metadata contains %d rows with duplicated normalized IDs; the first row is used.",
                       int(duplicate_mask.sum()))
    metadata = metadata.drop_duplicates("__normalized_id", keep="first").set_index("__normalized_id")

    groups: List[str] = []
    matches: List[bool] = []
    metadata_values: List[str] = []
    for genome in genome_names:
        key = normalize_id(genome)
        if key in metadata.index:
            matches.append(True)
            if group_column:
                raw = metadata.at[key, group_column]
                metadata_values.append(str(raw))
                groups.append(canonicalize_group(raw))
            else:
                metadata_values.append("")
                groups.append(default_group)
        else:
            matches.append(False)
            metadata_values.append("")
            groups.append("Unknown")

    audit["metadata_match"] = matches
    audit["raw_group_value"] = metadata_values
    audit["group"] = groups
    audit.to_csv(output_dir / "metadata_id_audit.tsv", sep="\t", index=False)
    LOGGER.info("Matched metadata for %d/%d genomes.", int(sum(matches)), len(matches))
    LOGGER.info("Group counts: %s", dict(Counter(groups)))
    return np.array(groups, dtype=object), id_column, group_column, audit


def read_alignment_headers(path: Path) -> List[str]:
    headers: List[str] = []
    with path.open("r", encoding="utf-8", errors="replace") as handle:
        for line in handle:
            if line.startswith(">"):
                headers.append(line[1:].strip().split()[0])
    return headers


def audit_alignment_ids(alignment_path: Optional[Path], genome_names: Sequence[str], output_dir: Path) -> None:
    if alignment_path is None:
        return
    if not alignment_path.exists():
        raise FileNotFoundError(f"Alignment file not found: {alignment_path}")
    headers = read_alignment_headers(alignment_path)
    if not headers:
        LOGGER.warning("No FASTA headers were found in the alignment; identifier audit skipped.")
        return
    rtab_norm = {normalize_id(x): x for x in genome_names}
    aln_norm = {normalize_id(x): x for x in headers}
    all_ids = sorted(set(rtab_norm) | set(aln_norm))
    rows = []
    for key in all_ids:
        rows.append({
            "normalized_id": key,
            "rtab_id": rtab_norm.get(key, ""),
            "alignment_id": aln_norm.get(key, ""),
            "in_rtab": key in rtab_norm,
            "in_alignment": key in aln_norm,
        })
    pd.DataFrame(rows).to_csv(output_dir / "alignment_id_audit.tsv", sep="\t", index=False)
    LOGGER.info("Alignment ID audit: %d headers; %d overlap with Rtab; %d only in Rtab; %d only in alignment.",
                len(headers), len(set(rtab_norm) & set(aln_norm)),
                len(set(rtab_norm) - set(aln_norm)), len(set(aln_norm) - set(rtab_norm)))


def compute_fluidity_matrix(binary_matrix: np.ndarray) -> Tuple[np.ndarray, np.ndarray]:
    """Return pairwise fluidity matrix and per-genome gene counts."""
    x = sparse.csr_matrix(binary_matrix.astype(np.int32, copy=False))
    gene_counts = np.asarray(x.sum(axis=0)).ravel().astype(np.float64)
    LOGGER.info("Computing pairwise shared-gene counts by sparse matrix multiplication.")
    shared = (x.T @ x).toarray().astype(np.float64, copy=False)
    denominator = gene_counts[:, None] + gene_counts[None, :]
    unique_total = denominator - 2.0 * shared
    with np.errstate(divide="ignore", invalid="ignore"):
        phi = np.divide(unique_total, denominator, out=np.zeros_like(unique_total), where=denominator > 0)
    phi = np.clip(phi, 0.0, 1.0)
    np.fill_diagonal(phi, 0.0)
    return phi, gene_counts


def upper_triangle_values(matrix: np.ndarray) -> np.ndarray:
    n = matrix.shape[0]
    return matrix[np.triu_indices(n, k=1)]


def subset_fluidity(matrix: np.ndarray, indices: np.ndarray) -> float:
    n = len(indices)
    if n < 2:
        return float("nan")
    sub = matrix[np.ix_(indices, indices)]
    return float(sub.sum() / (n * (n - 1)))


def jackknife_ci(matrix: np.ndarray) -> Tuple[float, float, float]:
    n = matrix.shape[0]
    if n < 3:
        return float("nan"), float("nan"), float("nan")
    total_upper = matrix.sum() / 2.0
    row_sums = matrix.sum(axis=1)
    denominator = (n - 1) * (n - 2) / 2.0
    leave_one_out = (total_upper - row_sums) / denominator
    loo_mean = float(leave_one_out.mean())
    se = math.sqrt((n - 1) / n * float(np.square(leave_one_out - loo_mean).sum()))
    mean_phi = float(matrix.sum() / (n * (n - 1)))
    low = max(0.0, mean_phi - 1.959963984540054 * se)
    high = min(1.0, mean_phi + 1.959963984540054 * se)
    return se, low, high


def genome_bootstrap(
    matrix: np.ndarray,
    n_replicates: int,
    batch_size: int,
    rng: np.random.Generator,
) -> np.ndarray:
    """Genome-resampling bootstrap for the order-2 U-statistic.

    A bootstrap sample is represented by multinomial genome multiplicities w.
    Because phi_ii = 0, the bootstrap estimate is w' H w / [n(n-1)].
    """
    n = matrix.shape[0]
    if n_replicates <= 0 or n < 2:
        return np.array([], dtype=float)
    probabilities = np.full(n, 1.0 / n)
    results = np.empty(n_replicates, dtype=np.float64)
    position = 0
    while position < n_replicates:
        b = min(batch_size, n_replicates - position)
        weights = rng.multinomial(n, probabilities, size=b).T.astype(np.float64)
        products = matrix @ weights
        values = np.sum(weights * products, axis=0) / (n * (n - 1))
        results[position:position + b] = values
        position += b
    return results


def summarize_fluidity(
    label: str,
    matrix: np.ndarray,
    n_gene_clusters: int,
    bootstrap_reps: int,
    bootstrap_batch: int,
    rng: np.random.Generator,
) -> Tuple[FluidityResult, np.ndarray]:
    values = upper_triangle_values(matrix)
    n = matrix.shape[0]
    mean_phi = float(values.mean())
    se, jk_low, jk_high = jackknife_ci(matrix)
    bootstrap_values = genome_bootstrap(matrix, bootstrap_reps, bootstrap_batch, rng)
    if bootstrap_values.size:
        boot_mean = float(bootstrap_values.mean())
        boot_low, boot_high = np.quantile(bootstrap_values, [0.025, 0.975])
    else:
        boot_mean = boot_low = boot_high = float("nan")
    result = FluidityResult(
        label=label,
        n_genomes=n,
        n_gene_clusters=n_gene_clusters,
        n_pairs=n * (n - 1) // 2,
        mean_phi=mean_phi,
        median_phi=float(np.median(values)),
        q1_phi=float(np.quantile(values, 0.25)),
        q3_phi=float(np.quantile(values, 0.75)),
        min_phi=float(values.min()),
        max_phi=float(values.max()),
        jackknife_se=se,
        jackknife_ci_low=jk_low,
        jackknife_ci_high=jk_high,
        bootstrap_mean=boot_mean,
        bootstrap_ci_low=float(boot_low),
        bootstrap_ci_high=float(boot_high),
    )
    return result, bootstrap_values


def result_to_dict(result: FluidityResult) -> Dict[str, object]:
    return {
        "dataset": result.label,
        "n_genomes": result.n_genomes,
        "n_gene_clusters": result.n_gene_clusters,
        "n_pairs": result.n_pairs,
        "mean_phi": result.mean_phi,
        "mean_percent_unshared": 100.0 * result.mean_phi,
        "median_phi": result.median_phi,
        "q1_phi": result.q1_phi,
        "q3_phi": result.q3_phi,
        "min_phi": result.min_phi,
        "max_phi": result.max_phi,
        "jackknife_se": result.jackknife_se,
        "jackknife_ci95_low": result.jackknife_ci_low,
        "jackknife_ci95_high": result.jackknife_ci_high,
        "bootstrap_mean": result.bootstrap_mean,
        "bootstrap_ci95_low": result.bootstrap_ci_low,
        "bootstrap_ci95_high": result.bootstrap_ci_high,
    }


def group_indices(groups: np.ndarray, min_size: int) -> Dict[str, np.ndarray]:
    output: Dict[str, np.ndarray] = {}
    for group in pd.unique(groups):
        idx = np.flatnonzero(groups == group)
        if len(idx) >= min_size and str(group) != "Unknown":
            output[str(group)] = idx
    return output


def pairwise_category_data(
    phi: np.ndarray,
    groups: np.ndarray,
    valid_groups: Dict[str, np.ndarray],
    max_plot_pairs: int,
    rng: np.random.Generator,
) -> Tuple[pd.DataFrame, Dict[str, np.ndarray]]:
    summaries: List[Dict[str, object]] = []
    plot_data: Dict[str, np.ndarray] = {}
    labels = list(valid_groups)
    for i, group_a in enumerate(labels):
        idx_a = valid_groups[group_a]
        for j in range(i, len(labels)):
            group_b = labels[j]
            idx_b = valid_groups[group_b]
            if i == j:
                sub = phi[np.ix_(idx_a, idx_a)]
                values = upper_triangle_values(sub)
                category = f"{group_a} within"
            else:
                values = phi[np.ix_(idx_a, idx_b)].ravel()
                category = f"{group_a} vs {group_b}"
            if values.size == 0:
                continue
            summaries.append({
                "category": category,
                "group_a": group_a,
                "group_b": group_b,
                "n_pairs": int(values.size),
                "mean_phi": float(values.mean()),
                "median_phi": float(np.median(values)),
                "q1_phi": float(np.quantile(values, 0.25)),
                "q3_phi": float(np.quantile(values, 0.75)),
                "min_phi": float(values.min()),
                "max_phi": float(values.max()),
            })
            if values.size > max_plot_pairs:
                selected = rng.choice(values.size, size=max_plot_pairs, replace=False)
                plot_data[category] = values[selected]
            else:
                plot_data[category] = values
    return pd.DataFrame(summaries), plot_data


def choose_comparison_groups(
    requested: Optional[Sequence[str]],
    valid_groups: Dict[str, np.ndarray],
) -> Optional[Tuple[str, str]]:
    if requested:
        a, b = requested
        if a not in valid_groups or b not in valid_groups:
            raise ValueError(f"Comparison groups must exactly match available groups: {list(valid_groups)}")
        return a, b
    clinical = next((g for g in valid_groups if "clinical" in g.lower()), None)
    environmental = next((g for g in valid_groups if "environment" in g.lower()), None)
    if clinical and environmental:
        return clinical, environmental
    if len(valid_groups) == 2:
        labels = list(valid_groups)
        return labels[0], labels[1]
    return None


def permutation_test(
    phi: np.ndarray,
    idx_a: np.ndarray,
    idx_b: np.ndarray,
    n_permutations: int,
    rng: np.random.Generator,
    batch_size: int = 100,
) -> Tuple[float, float, np.ndarray]:
    combined = np.concatenate([idx_a, idx_b])
    h = phi[np.ix_(combined, combined)]
    n_total = len(combined)
    n_a = len(idx_a)
    n_b = len(idx_b)
    observed = subset_fluidity(phi, idx_a) - subset_fluidity(phi, idx_b)
    if n_permutations <= 0:
        return observed, float("nan"), np.array([], dtype=float)

    diffs = np.empty(n_permutations, dtype=np.float64)
    pos = 0
    while pos < n_permutations:
        b = min(batch_size, n_permutations - pos)
        z = np.zeros((n_total, b), dtype=np.float64)
        for column in range(b):
            selected = rng.choice(n_total, size=n_a, replace=False)
            z[selected, column] = 1.0
        hz = h @ z
        a_values = np.sum(z * hz, axis=0) / (n_a * (n_a - 1))
        w = 1.0 - z
        hw = h @ w
        b_values = np.sum(w * hw, axis=0) / (n_b * (n_b - 1))
        diffs[pos:pos + b] = a_values - b_values
        pos += b
    p_value = (1.0 + float(np.sum(np.abs(diffs) >= abs(observed)))) / (n_permutations + 1.0)
    return observed, p_value, diffs


def matched_resampling(
    phi: np.ndarray,
    idx_a: np.ndarray,
    idx_b: np.ndarray,
    n_replicates: int,
    rng: np.random.Generator,
) -> pd.DataFrame:
    target = min(len(idx_a), len(idx_b))
    rows: List[Dict[str, float]] = []
    for replicate in range(1, n_replicates + 1):
        a = idx_a if len(idx_a) == target else rng.choice(idx_a, target, replace=False)
        b = idx_b if len(idx_b) == target else rng.choice(idx_b, target, replace=False)
        phi_a = subset_fluidity(phi, np.asarray(a))
        phi_b = subset_fluidity(phi, np.asarray(b))
        rows.append({
            "replicate": replicate,
            "target_n": target,
            "phi_group_a": phi_a,
            "phi_group_b": phi_b,
            "difference_a_minus_b": phi_a - phi_b,
        })
    return pd.DataFrame(rows)


def convergence_sizes(n: int) -> List[int]:
    candidates = [10, 20, 30, 50, 75, 100, 150, 250, 500, 750, 1000, 1250, n]
    return sorted(set(x for x in candidates if 2 <= x <= n))


def convergence_analysis(
    phi: np.ndarray,
    analysis_sets: Dict[str, np.ndarray],
    n_replicates: int,
    rng: np.random.Generator,
) -> pd.DataFrame:
    rows: List[Dict[str, object]] = []
    for label, indices in analysis_sets.items():
        n_total = len(indices)
        for size in convergence_sizes(n_total):
            reps = 1 if size == n_total else n_replicates
            estimates = np.empty(reps, dtype=float)
            for r in range(reps):
                selected = indices if size == n_total else rng.choice(indices, size=size, replace=False)
                estimates[r] = subset_fluidity(phi, np.asarray(selected))
            rows.append({
                "dataset": label,
                "sample_size": size,
                "replicates": reps,
                "mean_phi": float(estimates.mean()),
                "sd_phi": float(estimates.std(ddof=1)) if reps > 1 else 0.0,
                "ci95_low": float(np.quantile(estimates, 0.025)) if reps > 1 else float(estimates[0]),
                "ci95_high": float(np.quantile(estimates, 0.975)) if reps > 1 else float(estimates[0]),
            })
    return pd.DataFrame(rows)


def sensitivity_thresholds(n_genomes: int) -> List[Tuple[str, int]]:
    thresholds = [
        ("All clusters", 1),
        ("Prevalence >= 2 genomes", 2),
        ("Prevalence >= 5 genomes", 5),
        ("Prevalence >= 1%", max(2, math.ceil(0.01 * n_genomes))),
    ]
    output: List[Tuple[str, int]] = []
    seen = set()
    for label, threshold in thresholds:
        if threshold not in seen:
            output.append((label, threshold))
            seen.add(threshold)
    return output


def sensitivity_analysis(
    binary_matrix: np.ndarray,
    groups: np.ndarray,
    valid_groups: Dict[str, np.ndarray],
) -> pd.DataFrame:
    prevalence = binary_matrix.sum(axis=1)
    rows: List[Dict[str, object]] = []
    analysis_sets: Dict[str, np.ndarray] = {"All genomes": np.arange(binary_matrix.shape[1])}
    for label, idx in valid_groups.items():
        if len(idx) != binary_matrix.shape[1]:
            analysis_sets[label] = idx
    for scenario, threshold in sensitivity_thresholds(binary_matrix.shape[1]):
        keep = prevalence >= threshold
        filtered = binary_matrix[keep]
        LOGGER.info("Sensitivity scenario '%s': retaining %d/%d clusters.",
                    scenario, int(keep.sum()), len(keep))
        scenario_phi, _ = compute_fluidity_matrix(filtered)
        for label, idx in analysis_sets.items():
            if len(idx) < 2:
                continue
            value = subset_fluidity(scenario_phi, idx)
            rows.append({
                "scenario": scenario,
                "minimum_global_prevalence": threshold,
                "retained_gene_clusters": int(keep.sum()),
                "dataset": label,
                "n_genomes": len(idx),
                "mean_phi": value,
            })
    return pd.DataFrame(rows)


def save_figure(fig: plt.Figure, output_dir: Path, stem: str) -> None:
    fig.tight_layout()
    for extension in ("png", "pdf", "svg"):
        kwargs = {"bbox_inches": "tight"}
        if extension == "png":
            kwargs["dpi"] = 300
        fig.savefig(output_dir / f"{stem}.{extension}", **kwargs)
    plt.close(fig)


def plot_overall_distribution(values: np.ndarray, result: FluidityResult, output_dir: Path) -> None:
    fig, ax = plt.subplots(figsize=(8.0, 5.2))
    ax.hist(values, bins=80)
    ax.axvline(result.mean_phi, linestyle="--", linewidth=1.5, label=f"Mean = {result.mean_phi:.4f}")
    ax.axvline(result.median_phi, linestyle=":", linewidth=1.5, label=f"Median = {result.median_phi:.4f}")
    ax.set_xlabel("Pairwise genomic fluidity (phi_ij)")
    ax.set_ylabel("Number of genome pairs")
    ax.set_title("Distribution of pairwise gene-content fluidity")
    ax.legend(frameon=False)
    save_figure(fig, output_dir, "figure_01_pairwise_fluidity_distribution")


def plot_group_estimates(
    summary: pd.DataFrame,
    output_dir: Path,
    collection_label: Optional[str] = None,
) -> str:
    data = summary.copy()
    labels = data["dataset"].astype(str).tolist()
    if collection_label and len(labels) == 1 and labels[0] == "All genomes":
        labels = [f"{collection_label} collection"]
    means = data["mean_phi"].to_numpy(float)
    lows = data["jackknife_ci95_low"].to_numpy(float)
    highs = data["jackknife_ci95_high"].to_numpy(float)
    yerr = np.vstack([means - lows, highs - means])
    fig, ax = plt.subplots(figsize=(max(7.0, 1.1 * len(labels)), 5.2))
    positions = np.arange(len(labels))
    ax.bar(positions, means, yerr=yerr, capsize=4)
    ax.set_xticks(positions)
    ax.set_xticklabels(labels, rotation=25 if len(labels) > 1 else 0, ha="right" if len(labels) > 1 else "center")
    ax.set_ylabel("Mean genomic fluidity (phi)")
    if collection_label and len(labels) == 1:
        ax.set_title(f"Mean genomic fluidity of the {collection_label.lower()} collection")
    else:
        ax.set_title("Overall and group-specific genomic fluidity")
    ax.set_ylim(bottom=0)
    stem = (
        "figure_02_mean_fluidity_clinical_collection"
        if collection_label and collection_label.lower() == "clinical"
        else "figure_02_group_fluidity_estimates"
    )
    save_figure(fig, output_dir, stem)
    return stem


def plot_pairwise_categories(plot_data: Dict[str, np.ndarray], output_dir: Path) -> None:
    if len(plot_data) < 2:
        return
    labels = list(plot_data)
    values = [plot_data[label] for label in labels]
    fig, ax = plt.subplots(figsize=(max(8.0, 1.25 * len(labels)), 5.6))
    ax.violinplot(values, showmeans=True, showmedians=True, showextrema=False)
    ax.set_xticks(np.arange(1, len(labels) + 1))
    ax.set_xticklabels(labels, rotation=30, ha="right")
    ax.set_ylabel("Pairwise genomic fluidity (phi_ij)")
    ax.set_title("Within-group and between-group gene-content dissimilarity")
    save_figure(fig, output_dir, "figure_03_pairwise_group_categories")


def plot_convergence(convergence: pd.DataFrame, output_dir: Path) -> None:
    if convergence.empty:
        return
    fig, ax = plt.subplots(figsize=(8.2, 5.4))
    for label, data in convergence.groupby("dataset", sort=False):
        data = data.sort_values("sample_size")
        means = data["mean_phi"].to_numpy(float)
        lows = data["ci95_low"].to_numpy(float)
        highs = data["ci95_high"].to_numpy(float)
        yerr = np.vstack([means - lows, highs - means])
        ax.errorbar(data["sample_size"], means, yerr=yerr, marker="o", capsize=3, label=label)
    ax.set_xlabel("Number of genomes sampled")
    ax.set_ylabel("Mean genomic fluidity (phi)")
    ax.set_title("Sample-size convergence of genomic fluidity")
    ax.legend(frameon=False)
    save_figure(fig, output_dir, "figure_04_fluidity_convergence")


def plot_sensitivity(sensitivity: pd.DataFrame, output_dir: Path) -> None:
    if sensitivity.empty:
        return
    pivot = sensitivity.pivot(index="scenario", columns="dataset", values="mean_phi")
    scenario_order = sensitivity["scenario"].drop_duplicates().tolist()
    pivot = pivot.reindex(scenario_order)
    fig, ax = plt.subplots(figsize=(max(8.0, 1.6 * len(pivot)), 5.4))
    positions = np.arange(len(pivot))
    width = 0.8 / max(1, len(pivot.columns))
    for i, column in enumerate(pivot.columns):
        offset = (i - (len(pivot.columns) - 1) / 2) * width
        ax.bar(positions + offset, pivot[column].to_numpy(float), width=width, label=column)
    ax.set_xticks(positions)
    ax.set_xticklabels(pivot.index, rotation=25, ha="right")
    ax.set_ylabel("Mean genomic fluidity (phi)")
    ax.set_title("Sensitivity of fluidity to rare gene-cluster filtering")
    ax.set_ylim(bottom=0)
    ax.legend(frameon=False)
    save_figure(fig, output_dir, "figure_05_rare_gene_sensitivity")


def plot_gene_counts(gene_counts: np.ndarray, output_dir: Path) -> None:
    fig, ax = plt.subplots(figsize=(8.0, 5.2))
    ax.hist(gene_counts, bins=50)
    ax.axvline(np.median(gene_counts), linestyle="--", linewidth=1.5,
               label=f"Median = {np.median(gene_counts):.0f}")
    ax.set_xlabel("Gene clusters present per genome")
    ax.set_ylabel("Number of genomes")
    ax.set_title("Distribution of per-genome pangenome gene counts")
    ax.legend(frameon=False)
    save_figure(fig, output_dir, "figure_06_genome_gene_count_distribution")


def plot_matched_resampling(data: pd.DataFrame, group_a: str, group_b: str, output_dir: Path) -> None:
    if data.empty:
        return
    fig, ax = plt.subplots(figsize=(7.6, 5.2))
    ax.hist(data["difference_a_minus_b"], bins=50)
    ax.axvline(0.0, linestyle="--", linewidth=1.5)
    ax.set_xlabel(f"Matched fluidity difference: {group_a} - {group_b}")
    ax.set_ylabel("Resampling replicates")
    ax.set_title("Equal-sample-size fluidity comparison")
    save_figure(fig, output_dir, "figure_07_matched_group_difference")


def write_pairwise_values(
    phi: np.ndarray,
    genome_names: Sequence[str],
    groups: np.ndarray,
    output_dir: Path,
) -> None:
    n = len(genome_names)
    i_idx, j_idx = np.triu_indices(n, k=1)
    frame = pd.DataFrame({
        "genome_a": np.asarray(genome_names, dtype=object)[i_idx],
        "genome_b": np.asarray(genome_names, dtype=object)[j_idx],
        "group_a": groups[i_idx],
        "group_b": groups[j_idx],
        "phi_ij": phi[i_idx, j_idx],
    })
    frame.to_csv(output_dir / "pairwise_fluidity_values.tsv.gz", sep="\t", index=False, compression="gzip")


def first_stable_sample_size(convergence: pd.DataFrame, dataset: str, full_phi: float) -> Optional[int]:
    data = convergence[convergence["dataset"] == dataset].sort_values("sample_size")
    if data.empty or full_phi == 0:
        return None
    sizes = data["sample_size"].tolist()
    means = data["mean_phi"].to_numpy(float)
    for i, size in enumerate(sizes):
        relative_error = np.abs(means[i:] - full_phi) / full_phi
        if np.all(relative_error <= 0.025):
            return int(size)
    return None


def fmt(value: float, digits: int = 4) -> str:
    if value is None or not np.isfinite(value):
        return "not estimated"
    return f"{value:.{digits}f}"


def create_interpretation(
    organism: str,
    summary: pd.DataFrame,
    convergence: pd.DataFrame,
    sensitivity: pd.DataFrame,
    comparison: Optional[Tuple[str, str]],
    comparison_stats: Optional[Dict[str, float]],
    matched: Optional[pd.DataFrame],
    matrix_group_counts: Dict[str, int],
) -> Tuple[str, str, str]:
    overall = summary.iloc[0]
    n = int(overall["n_genomes"])
    clusters = int(overall["n_gene_clusters"])
    pairs = int(overall["n_pairs"])
    mean_phi = float(overall["mean_phi"])
    ci_low = float(overall["bootstrap_ci95_low"])
    ci_high = float(overall["bootstrap_ci95_high"])
    if not np.isfinite(ci_low):
        ci_low = float(overall["jackknife_ci95_low"])
        ci_high = float(overall["jackknife_ci95_high"])
    median = float(overall["median_phi"])

    methods = (
        f"Genomic fluidity was quantified from the binary Panaroo gene-presence/absence matrix for "
        f"{n:,} {organism} genomes. For each unordered pair of genomes k and l, pairwise fluidity was "
        f"calculated as (U_k + U_l)/(M_k + M_l), where U_k and U_l are gene clusters present exclusively "
        f"in one member of the pair and M_k and M_l are the total gene clusters present in each genome. "
        f"The collection-level estimate was the arithmetic mean across all {pairs:,} unique genome pairs. "
        f"Uncertainty was evaluated by genome-level nonparametric bootstrap resampling and a delete-one-genome "
        f"jackknife. Robustness to rare gene clusters was assessed after retaining clusters found in at least "
        f"two genomes, at least five genomes, and at least 1% of genomes. Random subsampling was used to examine "
        f"sample-size convergence."
    )

    results = (
        f"Across {n:,} genomes and {clusters:,} gene clusters, mean genomic fluidity was phi = {mean_phi:.4f} "
        f"(95% CI, {ci_low:.4f}-{ci_high:.4f}); the median pairwise value was {median:.4f}. "
        f"Thus, for an average pair of genomes, {mean_phi * 100:.2f}% of their combined gene-family occurrences "
        f"were restricted to one genome rather than shared by both."
    )

    all_sens = sensitivity[sensitivity["dataset"] == "All genomes"].copy()
    if not all_sens.empty:
        baseline = float(all_sens.iloc[0]["mean_phi"])
        strict = float(all_sens.iloc[-1]["mean_phi"])
        change = strict - baseline
        relative = abs(change) / baseline if baseline else float("nan")
        if np.isfinite(relative) and relative <= 0.10:
            results += (
                f" The estimate was comparatively robust to rare-cluster filtering: under the strictest "
                f"reported threshold, phi was {strict:.4f}, a change of {change:+.4f} from the unfiltered value."
            )
        else:
            results += (
                f" Rare gene clusters contributed materially to the estimate: under the strictest reported "
                f"threshold, phi changed to {strict:.4f} ({change:+.4f} relative to the unfiltered value)."
            )

    stable_n = first_stable_sample_size(convergence, "All genomes", mean_phi)
    if stable_n is not None and stable_n < n:
        results += (
            f" Random-subsampling analysis indicated that the point estimate remained within 2.5% of the "
            f"full-data estimate from approximately {stable_n:,} genomes onward."
        )

    if comparison and comparison_stats:
        group_a, group_b = comparison
        row_a = summary[summary["dataset"] == group_a]
        row_b = summary[summary["dataset"] == group_b]
        if not row_a.empty and not row_b.empty:
            a_phi = float(row_a.iloc[0]["mean_phi"])
            b_phi = float(row_b.iloc[0]["mean_phi"])
            diff = float(comparison_stats["observed_difference"])
            p = float(comparison_stats["permutation_p_value"])
            results += (
                f" Group-specific fluidity was {a_phi:.4f} for {group_a} and {b_phi:.4f} for {group_b} "
                f"(difference {group_a} - {group_b} = {diff:+.4f}; permutation P = {p:.4g})."
            )
            if matched is not None and not matched.empty:
                low, high = np.quantile(matched["difference_a_minus_b"], [0.025, 0.975])
                mean_diff = float(matched["difference_a_minus_b"].mean())
                results += (
                    f" Equal-size resampling produced a mean difference of {mean_diff:+.4f} "
                    f"(empirical 95% interval, {low:+.4f} to {high:+.4f})."
                )

    if len(matrix_group_counts) == 1:
        sole_group = next(iter(matrix_group_counts))
        if sole_group.lower() == "all":
            scope = (
                "No biological grouping variable was resolved from the metadata, so this estimate should be "
                "reported as collection-level gene-content heterogeneity. Source-specific conclusions require a "
                "validated grouping column."
            )
        else:
            scope = (
                f"Because the matrix contains a single analytical group ({sole_group}), this estimate should be "
                f"reported as gene-content heterogeneity within the {sole_group.lower()} collection rather than as a "
                f"clinical-versus-environmental comparison."
            )
    else:
        scope = (
            "Group differences describe gene-content structure in the sampled collection. They should not be "
            "interpreted as direct environmental or clinical selection effects without controlling for lineage, "
            "geography, assembly quality and sampling design."
        )

    interpretation = (
        f"The fluidity estimate quantifies average gene-family dissimilarity and complements, but does not replace, "
        f"pangenome accumulation or Heaps' law analysis. A value of {mean_phi:.4f} is a directly interpretable "
        f"effect size: an average genome pair differs in {mean_phi * 100:.2f}% of its combined gene-family "
        f"occurrences. The statistic alone does not identify the responsible genes and does not measure a rate of "
        f"horizontal gene transfer. Mechanistic interpretation therefore requires integration with functional "
        f"annotation, mobile-element analysis, population structure and phylogeny. {scope}"
    )
    return methods, results, interpretation


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def add_docx_table(document: Document, dataframe: pd.DataFrame, columns: Sequence[str], headers: Sequence[str]) -> None:
    table = document.add_table(rows=1, cols=len(columns))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    header_row = table.rows[0]
    set_repeat_table_header(header_row)
    for i, header in enumerate(headers):
        cell = header_row.cells[i]
        cell.text = header
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_shading(cell, "D9E1F2")
        for run in cell.paragraphs[0].runs:
            run.bold = True
            run.font.size = Pt(8.5)
    for _, row in dataframe.iterrows():
        cells = table.add_row().cells
        for i, column in enumerate(columns):
            value = row[column]
            if isinstance(value, (float, np.floating)):
                text = "NA" if not np.isfinite(value) else f"{value:.4f}"
            else:
                text = str(value)
            cells[i].text = text
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for run in cells[i].paragraphs[0].runs:
                run.font.size = Pt(8.5)


def create_docx_report(
    output_path: Path,
    organism: str,
    args: argparse.Namespace,
    summary: pd.DataFrame,
    category_summary: pd.DataFrame,
    convergence: pd.DataFrame,
    sensitivity: pd.DataFrame,
    methods: str,
    results: str,
    interpretation: str,
    comparison: Optional[Tuple[str, str]],
    comparison_stats: Optional[Dict[str, float]],
    output_dir: Path,
    collection_label: Optional[str] = None,
    figure2_stem: str = "figure_02_group_fluidity_estimates",
) -> None:
    if not DOCX_AVAILABLE:
        LOGGER.warning("python-docx is unavailable; DOCX report generation skipped.")
        return

    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    styles = document.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(10.5)
    for style_name in ("Title", "Heading 1", "Heading 2"):
        styles[style_name].font.name = "Arial"

    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("Pangenome Fluidity Analysis of ")
    species_run = title.add_run(organism)
    species_run.italic = True

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run("Reproducible analysis report generated from a Panaroo gene presence-absence matrix").italic = True
    subtitle.add_run(f"\nGenerated: {datetime.now().strftime('%Y-%m-%d %H:%M')}")

    document.add_heading("1. Analytical objective", level=1)
    document.add_paragraph(
        "This analysis quantifies within-collection gene-content variability using genomic fluidity. "
        "The statistic is the mean proportion of gene-family occurrences that are not shared by a pair of genomes."
    )

    document.add_heading("2. Input files and analytical role", level=1)
    inputs = pd.DataFrame([
        {"File": str(args.rtab), "Role": "Primary binary gene-cluster presence/absence matrix used for all fluidity calculations."},
        {"File": str(args.metadata) if args.metadata else "Not supplied", "Role": "Genome identifier matching and optional source/lineage grouping."},
        {"File": str(args.alignment) if args.alignment else "Not supplied", "Role": "Optional identifier audit only; not required for genomic fluidity."},
    ])
    add_docx_table(document, inputs, ["File", "Role"], ["Input", "Role in analysis"])

    document.add_heading("3. Definition and calculation", level=1)
    document.add_paragraph(
        "For genomes k and l, pairwise fluidity was calculated as: phi_kl = (U_k + U_l) / (M_k + M_l). "
        "U_k and U_l are the gene clusters restricted to one genome in the pair, while M_k and M_l are the total "
        "gene clusters present in the two genomes. Collection-level fluidity is the mean of phi_kl over all unordered "
        "genome pairs. With a binary matrix X, the numerator is the symmetric difference in gene-cluster content."
    )

    document.add_heading("4. Statistical workflow", level=1)
    for text in [
        "The complete, unfiltered Panaroo matrix was used for the primary estimate; core and accessory clusters were retained.",
        "Genome-level nonparametric bootstrap resampling generated percentile confidence intervals.",
        "A delete-one-genome jackknife provided an additional uncertainty estimate that does not treat genome pairs as independent observations.",
        "Rare-gene sensitivity analyses excluded singletons and progressively required global prevalence in at least 5 genomes or 1% of genomes.",
        "Random subsets of increasing size were used to examine convergence of the fluidity estimate.",
        "When two valid groups were available, within-group fluidities were compared by label permutation and equal-size resampling.",
    ]:
        document.add_paragraph(text, style="List Bullet")

    document.add_heading("5. Primary numerical results", level=1)
    display = summary.copy()
    display["95% bootstrap CI"] = display.apply(
        lambda r: f"{r['bootstrap_ci95_low']:.4f}-{r['bootstrap_ci95_high']:.4f}"
        if np.isfinite(r["bootstrap_ci95_low"]) else
        f"{r['jackknife_ci95_low']:.4f}-{r['jackknife_ci95_high']:.4f}", axis=1)
    add_docx_table(
        document,
        display,
        ["dataset", "n_genomes", "n_pairs", "mean_phi", "median_phi", "95% bootstrap CI"],
        ["Dataset", "Genomes", "Pairs", "Mean phi", "Median phi", "95% CI"],
    )

    document.add_heading("6. Publication-ready Methods paragraph", level=1)
    document.add_paragraph(methods)

    document.add_heading("7. Publication-ready Results paragraph", level=1)
    document.add_paragraph(results)

    document.add_heading("8. Scientific interpretation", level=1)
    document.add_paragraph(interpretation)

    figure2_caption = (
        f"Figure 2. Mean genomic fluidity of the {collection_label.lower()} collection with "
        "95% jackknife confidence intervals."
        if collection_label
        else "Figure 2. Mean genomic fluidity estimate(s) with 95% jackknife confidence intervals."
    )
    figure_specs = [
        ("figure_01_pairwise_fluidity_distribution.png", "Figure 1. Distribution of pairwise genomic fluidity values."),
        (f"{figure2_stem}.png", figure2_caption),
        ("figure_03_pairwise_group_categories.png", "Figure 3. Within-group and between-group pairwise gene-content dissimilarity."),
        ("figure_04_fluidity_convergence.png", "Figure 4. Sample-size convergence of the fluidity estimate."),
        ("figure_05_rare_gene_sensitivity.png", "Figure 5. Sensitivity of fluidity to the exclusion of rare gene clusters."),
        ("figure_06_genome_gene_count_distribution.png", "Figure 6. Distribution of gene-cluster counts per genome."),
        ("figure_07_matched_group_difference.png", "Figure 7. Equal-sample-size resampling distribution for the selected group comparison."),
    ]
    document.add_heading("9. Figures", level=1)
    for filename, caption in figure_specs:
        image = output_dir / filename
        if image.exists():
            document.add_picture(str(image), width=Inches(6.4))
            paragraph = document.add_paragraph(caption)
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.runs[0].italic = True
            paragraph.runs[0].font.size = Pt(9)

    if not category_summary.empty:
        document.add_heading("10. Pairwise category summary", level=1)
        add_docx_table(
            document,
            category_summary,
            ["category", "n_pairs", "mean_phi", "median_phi", "q1_phi", "q3_phi"],
            ["Pair category", "Pairs", "Mean phi", "Median phi", "Q1", "Q3"],
        )

    document.add_heading("11. Interpretation boundaries", level=1)
    for text in [
        "Fluidity is an effect-size measure of gene-content dissimilarity; it is not a direct estimate of horizontal gene-transfer rate.",
        "A larger value cannot be labelled universally high or low without an explicitly selected biological comparator.",
        "Differences between clinical and environmental collections may reflect lineage composition, geography, sequencing strategy, assembly quality or uneven sampling.",
        "The primary analysis should retain all gene clusters; rare-gene filtering is a sensitivity analysis rather than a replacement for the standard statistic.",
        "The core-genome alignment is not mathematically required for fluidity, but phylogenetic integration is needed to distinguish source effects from lineage effects.",
    ]:
        document.add_paragraph(text, style="List Bullet")

    document.add_heading("12. Reproducibility information", level=1)
    document.add_paragraph(
        f"Random seed: {args.seed}; bootstrap replicates: {args.bootstrap}; permutation replicates: "
        f"{args.permutations}; convergence replicates: {args.convergence_replicates}; matched resampling replicates: "
        f"{args.matched_replicates}. All numerical tables are available as tab-delimited files in the output directory."
    )

    document.add_heading("13. References", level=1)
    document.add_paragraph(
        "1. Kislyuk AO, Haegeman B, Bergman NH, Weitz JS. Genomic fluidity: an integrative view of gene diversity "
        "within microbial populations. BMC Genomics. 2011;12:32. doi:10.1186/1471-2164-12-32."
    )
    document.add_paragraph(
        "2. Tonkin-Hill G, MacAlasdair N, Ruis C, et al. Producing polished prokaryotic pangenomes with the Panaroo "
        "pipeline. Genome Biology. 2020;21:180. doi:10.1186/s13059-020-02090-4."
    )

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("Pangenome fluidity analysis report").font.size = Pt(8)
    document.save(output_path)


def write_publication_text(output_dir: Path, methods: str, results: str, interpretation: str) -> None:
    content = (
        "METHODS\n=======\n" + methods + "\n\n"
        "RESULTS\n=======\n" + results + "\n\n"
        "INTERPRETATION\n==============\n" + interpretation + "\n"
    )
    (output_dir / "publication_ready_methods_results.txt").write_text(content, encoding="utf-8")


def write_manifest(output_dir: Path, args: argparse.Namespace, id_column: Optional[str], group_column: Optional[str]) -> None:
    lines = [
        f"Analysis timestamp: {datetime.now().isoformat(timespec='seconds')}",
        f"Script: {Path(__file__).resolve()}",
        f"Rtab: {args.rtab.resolve()}",
        f"Metadata: {args.metadata.resolve() if args.metadata else 'Not supplied'}",
        f"Alignment: {args.alignment.resolve() if args.alignment else 'Not supplied'}",
        f"Metadata ID column: {id_column or 'Not resolved'}",
        f"Metadata group column: {group_column or 'Not resolved'}",
        f"Organism: {args.organism}",
        f"Random seed: {args.seed}",
        f"Bootstrap replicates: {args.bootstrap}",
        f"Permutation replicates: {args.permutations}",
        f"Convergence replicates: {args.convergence_replicates}",
        f"Matched resampling replicates: {args.matched_replicates}",
        f"Python: {sys.version.replace(chr(10), ' ')}",
        f"NumPy: {np.__version__}",
        f"pandas: {pd.__version__}",
    ]
    (output_dir / "analysis_manifest.txt").write_text("\n".join(lines) + "\n", encoding="utf-8")


def main() -> int:
    args = parse_args()
    configure_logging(args.output)
    LOGGER.info("Starting pangenome fluidity analysis.")
    rng = np.random.default_rng(args.seed)

    try:
        binary_matrix, gene_names, genome_names = read_rtab(args.rtab)
        groups, resolved_id_col, resolved_group_col, metadata_audit = map_metadata_groups(
            args.metadata, genome_names, args.id_column, args.group_column,
            args.default_group, args.output,
        )
        audit_alignment_ids(args.alignment, genome_names, args.output)

        phi, gene_counts = compute_fluidity_matrix(binary_matrix)
        genome_count_frame = pd.DataFrame({
            "genome": genome_names,
            "group": groups,
            "gene_clusters_present": gene_counts.astype(int),
        })
        genome_count_frame.to_csv(args.output / "genome_gene_counts.tsv", sep="\t", index=False)

        valid_groups = group_indices(groups, args.min_group_size)
        analysis_sets: Dict[str, np.ndarray] = {"All genomes": np.arange(len(genome_names))}
        for label, idx in valid_groups.items():
            if len(idx) != len(genome_names):
                analysis_sets[label] = idx

        summary_results: List[FluidityResult] = []
        bootstrap_store: Dict[str, np.ndarray] = {}
        for label, idx in analysis_sets.items():
            LOGGER.info("Summarizing fluidity for %s (n=%d).", label, len(idx))
            sub = phi[np.ix_(idx, idx)]
            detected_clusters = int(np.any(binary_matrix[:, idx] > 0, axis=1).sum())
            result, bootstrap_values = summarize_fluidity(
                label, sub, detected_clusters, args.bootstrap,
                args.bootstrap_batch, rng,
            )
            summary_results.append(result)
            bootstrap_store[label] = bootstrap_values

        summary = pd.DataFrame([result_to_dict(x) for x in summary_results])
        summary.to_csv(args.output / "fluidity_summary.tsv", sep="\t", index=False, float_format="%.10g")

        overall_values = upper_triangle_values(phi)
        if args.save_pairwise:
            LOGGER.info("Writing all pairwise fluidity values.")
            write_pairwise_values(phi, genome_names, groups, args.output)

        category_summary, category_plot_data = pairwise_category_data(
            phi, groups, valid_groups, args.max_plot_pairs, rng
        )
        category_summary.to_csv(args.output / "pairwise_category_summary.tsv", sep="\t", index=False,
                                float_format="%.10g")

        comparison = choose_comparison_groups(args.compare_groups, valid_groups)
        comparison_stats: Optional[Dict[str, float]] = None
        matched: Optional[pd.DataFrame] = None
        if comparison:
            group_a, group_b = comparison
            idx_a = valid_groups[group_a]
            idx_b = valid_groups[group_b]
            LOGGER.info("Comparing '%s' (n=%d) against '%s' (n=%d).", group_a, len(idx_a), group_b, len(idx_b))
            observed, p_value, permutation_distribution = permutation_test(
                phi, idx_a, idx_b, args.permutations, rng
            )
            comparison_stats = {
                "group_a": group_a,
                "group_b": group_b,
                "n_group_a": len(idx_a),
                "n_group_b": len(idx_b),
                "observed_difference": observed,
                "permutation_p_value": p_value,
                "permutations": args.permutations,
            }
            pd.DataFrame([comparison_stats]).to_csv(
                args.output / "group_permutation_test.tsv", sep="\t", index=False, float_format="%.10g"
            )
            if permutation_distribution.size:
                pd.DataFrame({"permuted_difference": permutation_distribution}).to_csv(
                    args.output / "group_permutation_distribution.tsv.gz", sep="\t", index=False,
                    compression="gzip", float_format="%.10g"
                )
            if args.matched_replicates > 0:
                matched = matched_resampling(phi, idx_a, idx_b, args.matched_replicates, rng)
                matched.to_csv(args.output / "matched_group_resampling.tsv.gz", sep="\t", index=False,
                               compression="gzip", float_format="%.10g")
        else:
            LOGGER.info("No unambiguous two-group comparison was available; permutation and matched analyses skipped.")

        convergence_sets = {"All genomes": np.arange(len(genome_names))}
        for label, idx in valid_groups.items():
            if len(idx) != len(genome_names) and len(convergence_sets) < 4:
                convergence_sets[label] = idx
        convergence = convergence_analysis(phi, convergence_sets, args.convergence_replicates, rng)
        convergence.to_csv(args.output / "fluidity_convergence.tsv", sep="\t", index=False,
                           float_format="%.10g")

        sensitivity = sensitivity_analysis(binary_matrix, groups, valid_groups)
        sensitivity.to_csv(args.output / "rare_gene_sensitivity.tsv", sep="\t", index=False,
                           float_format="%.10g")

        plot_overall_distribution(overall_values, summary_results[0], args.output)
        nonunknown_counts = {
            str(k): int(v) for k, v in Counter(groups).items() if str(k) != "Unknown"
        }
        collection_label = (
            next(iter(nonunknown_counts))
            if len(nonunknown_counts) == 1 and next(iter(nonunknown_counts.values())) == len(genome_names)
            else None
        )
        figure2_stem = plot_group_estimates(summary, args.output, collection_label)
        plot_pairwise_categories(category_plot_data, args.output)
        plot_convergence(convergence, args.output)
        plot_sensitivity(sensitivity, args.output)
        plot_gene_counts(gene_counts, args.output)
        if matched is not None and comparison:
            plot_matched_resampling(matched, comparison[0], comparison[1], args.output)

        group_counts = nonunknown_counts
        methods, results, interpretation = create_interpretation(
            args.organism, summary, convergence, sensitivity, comparison,
            comparison_stats, matched, group_counts,
        )
        write_publication_text(args.output, methods, results, interpretation)
        create_docx_report(
            args.output / "Pangenome_Fluidity_Results_Report.docx",
            args.organism, args, summary, category_summary, convergence,
            sensitivity, methods, results, interpretation, comparison,
            comparison_stats, args.output, collection_label, figure2_stem,
        )
        write_manifest(args.output, args, resolved_id_col, resolved_group_col)

        LOGGER.info("Analysis completed successfully. Primary fluidity: phi = %.6f.", summary_results[0].mean_phi)
        LOGGER.info("Results directory: %s", args.output.resolve())
        return 0
    except Exception as exc:
        LOGGER.exception("Analysis failed: %s", exc)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
