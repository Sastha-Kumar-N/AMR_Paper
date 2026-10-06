"""Finalize the evidence audit, workflow workbook and reproducibility archive."""
from pathlib import Path
import sys,json,re,hashlib,shutil,zipfile,gzip,importlib.metadata as md
import pandas as pd
OUT=Path(__file__).resolve().parents[1];BASE=OUT.parent;ROOT=BASE.parent;D=OUT/'source_data'
sys.path.append(str(ROOT/'18.amrfinderplus/amrfinderplus_interpretation/envs/amrinterp/lib/python3.13/site-packages'))
sys.path.append(str(ROOT/'17.clinical_comparative_genomics/map_genome/map_pan_3/map_pan_3/.venv/lib/python3.13/site-packages'))
from openpyxl import Workbook,load_workbook
from openpyxl.styles import Font,PatternFill,Alignment
from openpyxl.utils import get_column_letter
A=json.loads((D/'final_numerical_validation.json').read_text());L=json.loads((D/'rendered_layout_validation.json').read_text());M=json.loads((D/'manuscript_content.json').read_text());W=json.loads((BASE/'workflow_catalogue.json').read_text())
assert A['all_passed'] and all(x['all_passed'] for x in L)
# Clear final PDF names; the rendering work area retains audit previews separately.
for name in ['P_aeruginosa_Research_Article','Supplementary_Information']:shutil.copyfile(OUT/'rendered/final'/f'{name}.pdf',OUT/f'{name}.pdf')

extra=[
 {'step':'W40','experiment':'Occupancy-weighted pangenome content','population':'1,595 whole; 1,350 clinical; 99 environmental-labelled assemblies','what_done':'Recomputed cluster prevalence bands and pooled binary genome–cluster presences from original matrices.','result':'Whole: 22,567 clusters occur in ≤5 assemblies; 52.48% of distinct clusters but 0.44% of 9,700,435 presences. ≥99% core contributes 74.74% of presences.','interpreted_output':'The catalogue of alternative genes is much larger than its contribution to pooled genome content.','finding':'Rare-repertoire richness and per-genome content are different measures.','status':'Verified new analysis','limitation':'Presence weighting is not nucleotide mass or copy number; cluster calls remain annotation-dependent.','method':'Integer-safe occupancy and pooled presence calculations.','source':'article_20260911/source_data/occupancy_weighted_architecture.tsv; per_genome_frequency_classes.tsv'},
 {'step':'W41','experiment':'Exact finite-collection rarefaction and power-law modelling','population':'Original whole, clinical-labelled and environmental-labelled matrices','what_done':'Computed hypergeometric expected pan/core/discovery curves; fitted accumulation and discovery models independently.','result':'Whole γ=0.290002, κ=5070.301; clinical γ=0.288462, κ=5101.381. Final expected gains 7.707 and 8.668.','interpreted_output':'Accumulation remains unsaturated within the sampled collection.','finding':'Exact expectation removes random-order Monte Carlo variation; endpoint gain equals singletons/N.','status':'Verified new primary manuscript model','limitation':'No extrapolation or population CI; historical W14/W15 fits remain different estimators.','method':'Non-negative nonlinear least squares; independent exhaustive-subset and hypergeometric validation.','source':'article_20260911/source_data/*exact_rarefaction.tsv; mathematics_validation.json'},
 {'step':'W42','experiment':'Known repeated-BioSample sensitivity','population':'Retained 1,592 whole and 1,347 clinical-labelled assemblies','what_done':'Excluded three specified GCA records sharing BioSamples and primary chromosomes with retained GCF records; held original clusters fixed.','result':'Whole γ=0.290250; clinical γ=0.288760; clinical fluidity=0.107298. AMR median 10 and IQR 8–16 unchanged; 928 symbols retained.','interpreted_output':'These three known duplicate biological representations do not drive the primary descriptive patterns.','finding':'A bounded biological-representation sensitivity supports stability of the estimates.','status':'Verified new sensitivity','limitation':'Not a Panaroo rerun or comprehensive correction for clonality.','method':'Column exclusion, occupancy recalculation, exact rarefaction, integer-safe pairwise fluidity and AMR symbol collapse.','source':'article_20260911/source_data/additional_validated_results.json'},
 {'step':'W43','experiment':'Manuscript, figure and numerical reconciliation','population':'Final article and supplementary outputs','what_done':'Wrote/reviewed/revised article; verified references; generated five main and two supplementary figures; embedded eight editable equations; rendered DOCX.','result':'141 numerical checks passed; 114 publication metrics traced; 2,772 original source files unchanged; seven figures reproduce identically; no figure/text overlap or blank pages.','interpreted_output':'A coherent, auditable manuscript package is available for author completion.','finding':'Known erroneous or incomplete branches are excluded from primary claims.','status':'Completed with explicit submission requirements','limitation':'Funding, author approval, deposition links and unrecoverable runtime/database provenance require resolution before submission.','method':'SHA-256, direct matrix/statistical checks, Office Math XML and rendered-layout validation.','source':'13_final_consistency_audit.md; article_20260911/source_data/*validation.json'}]
work=[dict(x) for x in W]+extra
for r in work:
 if r['step'] in ['W14','W15']:
  r['status']='Verified historical Monte Carlo analysis'
  r['limitation']='The current article uses W41 exact finite-collection curves. '+r['limitation']
book=Workbook();book.remove(book.active)
def sheet(name,headers,rows,desc,widths=None):
 ws=book.create_sheet(name);ws.sheet_view.showGridLines=False;ws.merge_cells(start_row=1,start_column=1,end_row=1,end_column=len(headers));ws.cell(1,1,desc);ws.cell(1,1).font=Font(name='Calibri',size=15,bold=True,color='17344B');ws.cell(1,1).alignment=Alignment(wrap_text=True,vertical='center');ws.row_dimensions[1].height=44
 ws.append([]);ws.append(headers);ws.freeze_panes='D4' if name=='Workflow' else 'B4';ws.auto_filter.ref=f'A3:{get_column_letter(len(headers))}{len(rows)+3}'
 for c in ws[3]:c.font=Font(name='Calibri',size=11,bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='17344B');c.alignment=Alignment(wrap_text=True)
 ws.row_dimensions[3].height=35
 for i,row in enumerate(rows,4):
  ws.append([None if pd.isna(x) else x for x in row]);ws.row_dimensions[i].height=100 if name=='Workflow' else 52
  for c in ws[i]:
   c.font=Font(name='Calibri',size=11);c.alignment=Alignment(wrap_text=True,vertical='top')
   if isinstance(c.value,str):c.data_type='s'
   if i%2==0:c.fill=PatternFill('solid',fgColor='EFF5F7')
 for j in range(1,len(headers)+1):ws.column_dimensions[get_column_letter(j)].width=widths[j-1] if widths else (22 if j==1 else 42)
 return ws
keys=['step','experiment','population','what_done','result','interpreted_output','finding','status','limitation','method','source']
sheet('Workflow',['Step','Analysis / experiment','Population / input','What we did','Verified result','Interpreted output','Distinctive finding','Evidence status','Limits / missing work','Method','Exact source'],[[r.get(k,'') for k in keys] for r in work],'Complete workflow: original analyses plus verified article extensions. W41 is the primary accumulation model.',[10,35,38,60,65,60,55,30,60,50,75])
metrics=pd.read_csv(D/'publication_metric_registry.tsv',sep='\t').fillna('')
sheet('Publication_numbers',metrics.columns.to_list(),metrics.values.tolist(),'Numerical reference for the revised article: 114 source-traced metric rows.')
figs=json.loads((OUT/'figures/figure_manifest.json').read_text())
sheet('Figures',['Figure','File','Source','Evidence unit'],[[f['figure'],f['png'],'; '.join(f.get('source_tables',[])) or f.get('original_source',''),'See figure legend and source table; no biological activity inferred from prediction'] for f in figs],'Five main figures and two supplementary figures; all are embedded and independently auditable.')
gates=[['Provenance','Recover original Panaroo runtime and missing annotation/database releases, or document an appropriately versioned rerun.','Required before submission'],['Data/code availability','Deposit the local prepared package and supply real repository URLs/accessions/DOIs.','Required before submission'],['Author declarations','Confirm author order, contributions, funding, competing interests and approval of new analyses/text.','Required before submission'],['Habitat inference','Recurate metadata and adjust for population structure before ecological claims.','Required only for stronger habitat claims'],['Phylogeny / BGC','Resolve bootstrap convergence and BGC region exports before restoring their excluded claims.','Required only for those claims'],['Function / phenotype','Add AST or experimental functional testing before resistance-phenotype, AMP activity or vaccine claims.','Required only for those claims']]
sheet('Submission_items',['Area','Action','Status'],gates,'The manuscript is submission-formatted; unresolved evidence and author fields are not fabricated.',[30,95,40])
wpath=OUT/'P_aeruginosa_Article_Workflow_20260911.xlsx';book.save(wpath)
checkbook=load_workbook(wpath,read_only=True,data_only=False);assert checkbook['Workflow'].max_row==46;assert checkbook['Publication_numbers'].max_row==117
for ws in checkbook:
 for row in ws.iter_rows():assert all(c.data_type!='f' for c in row)
checkbook.close();print('Workflow workbook: 43 stages, 4 sheets.',flush=True)

# Portable figure requirements and original matrix copies (compressed, never modified).
vers={}
for pkg in ['numpy','pandas','scipy','matplotlib','python-docx','openpyxl','lxml','Pillow']:
 try:vers[pkg]=md.version(pkg)
 except md.PackageNotFoundError:vers[pkg]='not recovered'
(OUT/'scripts/requirements_recorded.txt').write_text('\n'.join(k+'=='+v for k,v in vers.items() if v!='not recovered')+'\n')
(D/'recorded_package_versions.json').write_text(json.dumps(vers,indent=2))
inp=OUT/'reproducibility_inputs';inp.mkdir(exist_ok=True);compressed=[]
for pop in ['whole','clinical','environmental']:
 src=ROOT/f'14.pangenome_panaroo/results/{pop}/gene_presence_absence.Rtab';dest=inp/f'{pop}_gene_presence_absence.Rtab.gz'
 if not dest.exists():
  with src.open('rb') as a,gzip.open(dest,'wb',compresslevel=6) as b:shutil.copyfileobj(a,b,2**20)
 h=hashlib.sha256()
 with gzip.open(dest,'rb') as f:
  for x in iter(lambda:f.read(2**20),b''):h.update(x)
 expected=json.loads((BASE/'core_source_checksums.json').read_text())[str(src.relative_to(ROOT))]['sha256'];assert h.hexdigest()==expected
 compressed.append({'file':dest.name,'original_path':str(src.relative_to(ROOT)),'uncompressed_sha256':expected})
(inp/'manifest.json').write_text(json.dumps(compressed,indent=2))

audit=f'''# 13 — Final consistency audit

Completed 11 September 2026 against `03_MASTER_STUDY_EVIDENCE.md`, original-result checksum registries and the revised article source model.

## Outcome

**Numerical and rendered-document checks passed. Submission readiness remains conditional on the explicitly marked author, deposition and provenance requirements.** No missing analysis, version, funding award, approval or repository identifier has been invented.

| Check | Result | Evidence |
|---|---|---|
| Automated numerical / traceability checks | {len(A['checks'])} passed | article_20260911/source_data/final_numerical_validation.json |
| Publication metric rows | {A['publication_metrics']} verified | publication_metric_registry.tsv; appended to master evidence |
| Numeric-token review index | {A['numeric_tokens_review_index']} occurrences | all_manuscript_numeric_tokens.tsv; includes citations, identifiers and method settings |
| Original source integrity | {A['original_sources_verified']:,} files unchanged; {A['source_bytes_verified']:,} bytes checked | original_source_integrity.tsv; relocations explicitly recorded |
| Exact rarefaction | Exhaustive small-subset, direct hypergeometric, monotonicity and endpoint checks passed | mathematics_validation.json |
| Table 1 | All three population counts and category partitions match original matrices | final_numerical_validation.json |
| Figure reproducibility | Seven PNG figures reproduced byte-identically from packaged source tables | figure_portability_validation.json |
| Word structure | Five main + two supplementary inline figures; eight editable Office Math equations; zero floating anchors | docx_structure_validation.json |
| Final render | 35 article pages; 7 supplementary pages; zero blank pages, out-of-page text or image/text overlaps; all captions on figure pages | rendered_layout_validation.json |
| Supplementary data | Six reopened workbooks; full TSV equivalents; no unintended Excel formulas | supplementary_data_manifest.json |
| Updated workflow | 43 stages and four sheets; original workflow workbook preserved | P_aeruginosa_Article_Workflow_20260911.xlsx |
| References | 25 verified primary/software references; all cited; adjacent references separated by commas in superscripts | verified_references.json |
| Journal-style structure | 12-word title; 143-word unreferenced abstract; ~2,995 words in Introduction/Results/Discussion/Conclusion; Methods separate | article_structure_check.json; manuscript_content.json |

## Population and unit reconciliation

The traceable input is 3,765 files, leading to 1,630 exact-hash representatives, 1,619 QUAST passes, 1,604 CheckM2 passes and 1,595 GTDB-defined study assemblies. The historical source populations are 1,350 clinical-labelled, 99 environmental-labelled and 146 other assemblies. Sensitivity populations of 1,592/1,347 are explicitly separated. The original whole/clinical/environmental cluster totals are 43,001/40,811/18,475; alignment reruns are not substituted.

Strict core and ≥99% core are distinguished (whole 2,352 versus 4,551). Four prevalence-band counts sum to their corresponding repertoire totals; singleton/doubleton counts are nested. Rare-cluster distinct fractions are separated from pooled binary-presence fractions. Standalone AMR uses 929 exact symbols, including mutation/disruption symbols; the separate 204-family ARG analysis has a different unit and database history. AMP occurrences, unique sequences and similarity families remain separate; CAZyme domain tokens, gene–family pairs and genes remain separate.

## Explicitly reconciled revisions

1. The primary article uses the new exact finite-collection accumulation fits: whole γ = 0.290002 and clinical γ = 0.288462, with final expected gains 7.707/8.668. Historical permutation fits and order-specific terminal gains are retained only as labelled historical results. A stale F02 row in the planning/claim ledger was corrected to the chosen exact estimator during this audit.
2. The jackknife limits in the original corrected fluidity file use the precise normal quantile; one independent audit used 1.96. The approximately 4 × 10^−8 endpoint difference disappears at the reported precision and is explained in the master evidence. Bootstrap and jackknife intervals are not interchanged.
3. The later 1,347-assembly geographic branch has 601 coordinate-bearing records, while historical clinical metadata have 603/1,350. No coordinate number appears as an AMR hotspot count.
4. Misaligned source comparisons, overflowing distances, non-converged phylogenetic conclusions, inconsistent BGC loci and absent ML/vaccine validations were excluded or explicitly qualified.

## Peer-review and layout changes

The revised article adds jackknife uncertainty and the CAZyme distinct-gene IQR. Figure 2's initial heading overlap was corrected. Two blank pages and an invalid footer field were fixed. Mathematical accents and subscripts are editable; adjacent citations render as 1,2 rather than an ambiguous 12. The supplementary Environmental column heading was abbreviated to avoid mid-word wrapping. Figures were visually inspected, and rendered bounding boxes were checked for image/text intersections and page overflow. Layout verification used LibreOffice 25.2.3.2; a different word processor may paginate differently.

## Remaining submission requirements

- Recover or otherwise resolve the historical Panaroo runtime and missing CheckM2, Prokka, AMP-reference and dbCAN database releases; the exact placeholder is retained wherever unavailable.
- Supply actual public data/code deposition identifiers. A local reproducibility package is prepared; no external publication was performed.
- Confirm author list/contributions, approval of new analyses and revised text, funding and competing-interest statements.
- Recurate metadata and address lineage structure before stronger habitat conclusions. Resolve phylogeny/BGC problems before restoring their excluded claims. These analyses are not presented as completed results.

## Deliverables

All 13 requested reconstruction Markdown files are present. The article package contains the revised DOCX/PDF, Supplementary Information DOCX/PDF, six source-data workbooks and TSVs, updated workflow workbook, figures, reusable Python scripts, reference ledger, compressed copies of the original three occupancy matrices, and integrity/validation reports. The original project was treated as read-only. External directory relocation is documented separately and was not performed by the agent.
'''
(BASE/'13_final_consistency_audit.md').write_text(audit)

readme='''# Pseudomonas aeruginosa article package

Start with **P_aeruginosa_Research_Article.docx** (editable article) or its PDF. **Supplementary_Information.docx/PDF** contains the corrected existing fluidity figure, exploratory source associations and diagnostic methods. **P_aeruginosa_Article_Workflow_20260911.xlsx** describes all 43 workflow stages, outputs, interpretation and limitations.

The main article is formatted in a Communications Biology/Nature-family research-article style, with a short abstract, Introduction, Results, Discussion, Conclusion, Methods, declarations and numbered references. Journal acceptance or Q1/T1 placement is not guaranteed. The explicit author and version fields must be resolved before submission; see the workflow Submission_items sheet and the final consistency audit.

## Data and figures

- `supplementary_data/`: six Excel workbooks and lossless TSV copies; the numbering is used in the manuscript.
- `figures/`: five main figures and two supplementary figures; PNG for Word, and PDF/SVG for all new plots. Supplementary Fig. 1 is a verified existing image reused without modification.
- `source_data/`: exact curves, per-assembly summaries, audit snapshots, all figure source data, manuscript content, references and validation reports.
- `reproducibility_inputs/`: compressed byte-verified copies of the three original occupancy matrices, with original paths and uncompressed hashes.
- `scripts/`: reusable Python scripts. No AI-generated scientific raster figures are used; all new plots derive from numeric results.

## Reproduce figures independently

Keep the `scripts/`, `source_data/` and `figures/` directories together. In an environment with NumPy, pandas and Matplotlib, run:

```bash
PYTHONDONTWRITEBYTECODE=1 python3 -B scripts/03_make_figures.py
```

This script reads the included source tables and audit snapshot, including a preserved copy of the original corrected fluidity image. It does not need the external original project directories. All seven PNGs were reproduced identically with this input mode. Record exact package versions using `scripts/requirements_recorded.txt`; use an isolated environment if installing dependencies.

## Recompute and audit in the original project

The exact-analysis script `01_validate_and_extend.py`, supplement builder `06_make_supplements.py` and final numerical audit `08_final_audit.py` resolve original inputs relative to `<PROJECT_FOLDER>/manuscript_reconstruction/article_20260911/`. They require the preserved original result files at the paths in the source registries. The original audit scripts and their outputs are supplied alongside the reconstruction evidence. Original computational databases/software not recovered from logs are explicitly marked; this package does not claim that every upstream pipeline can be regenerated from an unavailable environment.

`04_validate_mathematics.py` checks exact expectations against exhaustive subsets and binomial calculations. `05_author_manuscript.py` retains the structured text, evidence map and internal revision. `07_make_docx.py` creates native Word equations and inline figures. These scripts write only reconstruction outputs. `02_build_evidence.py` deliberately refuses to overwrite completed evidence files.

For local PDF layout verification, LibreOffice can convert the generated DOCX files; Poppler's pdftotext/pdftohtml can inspect page text and image bounds. A temporary LibreOffice profile was used. No documents were emailed, submitted or deposited to an external service.

## Evidence hierarchy

The source outputs and checksum registries underpin `03_MASTER_STUDY_EVIDENCE.md`; `10_manuscript_evidence_map.md` maps Results statements to sources; `11_internal_peer_review.md` records the critique and revisions; `13_final_consistency_audit.md` records final validation and unresolved submission items. Historical stochastic fits and original source labels remain documented, distinct from the article's new exact models and qualified interpretations.
'''
(OUT/'README.md').write_text(readme)
for i in range(1,14):assert list(BASE.glob(f'{i:02d}_*.md')),i
# Include the evidence and compact audit outputs in their original relative hierarchy.
selected=[]
for f in OUT.rglob('*'):
 if not f.is_file():continue
 rel=f.relative_to(OUT)
 if any(x in rel.parts for x in ['rendered','matplotlib_cache','__pycache__']):continue
 if f.name.startswith('QA_') or f.suffix=='.zip' or f.name=='delivery_manifest.json':continue
 selected.append(f)
for f in BASE.iterdir():
 if not f.is_file():continue
 if (re.match(r'^\d{2}_.*\.md$',f.name) or f.name.startswith(('verified_','audit_')) or f.name.endswith('_source_checksums.json') or f.name in ['study_workflow_data.py','workflow_catalogue.json','quantitative_result_catalogue.json','source_path_relocations.tsv','metadata_manual_review_priority.tsv']):selected.append(f)
selected=sorted(set(selected));manifest=[]
for f in selected:manifest.append({'path':str(f.relative_to(BASE)),'bytes':f.stat().st_size,'sha256':hashlib.sha256(f.read_bytes()).hexdigest()})
(OUT/'delivery_manifest.json').write_text(json.dumps(manifest,indent=2))
archive=OUT/'P_aeruginosa_Article_Package.zip'
with zipfile.ZipFile(archive,'w',compression=zipfile.ZIP_DEFLATED,compresslevel=6) as z:
 for f in selected:z.write(f,'manuscript_reconstruction/'+str(f.relative_to(BASE)))
 z.write(OUT/'delivery_manifest.json','manuscript_reconstruction/article_20260911/delivery_manifest.json')
with zipfile.ZipFile(archive) as z:assert z.testzip() is None
print('Package',archive.name,'files',len(manifest)+1,'bytes',archive.stat().st_size,flush=True)
