"""Create the clean Word master using the authoritative bundled Python runtime.
Native OMML represents the two display equations; body and tables follow the MD master.
No original file is changed. Figures remain separate vector source-linked artifacts.
"""
from pathlib import Path
import re,json
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.opc.constants import RELATIONSHIP_TYPE as RT
R=Path(__file__).resolve().parents[2];O=R/'publication_manuscript'
doc=Document();sec=doc.sections[0]
sec.page_width=Inches(8.5);sec.page_height=Inches(11)
sec.top_margin=sec.bottom_margin=Inches(.78);sec.left_margin=sec.right_margin=Inches(.8)
sec.header_distance=sec.footer_distance=Inches(.32)
styles=doc.styles
for n in ['Normal','Body Text']:
 st=styles[n];st.font.name='Times New Roman';st.font.size=Pt(11);st.font.color.rgb=RGBColor.from_string('000000');st.paragraph_format.line_spacing=1.12;st.paragraph_format.space_after=Pt(6)
for n,size in [('Title',17),('Heading 1',13),('Heading 2',11.5),('Heading 3',11)]:
 st=styles[n];st.font.name='Times New Roman';st.font.size=Pt(size);st.font.bold=True;st.font.color.rgb=RGBColor.from_string('000000');st.paragraph_format.space_before=Pt(11);st.paragraph_format.space_after=Pt(5);st.paragraph_format.keep_with_next=True
styles['Title'].paragraph_format.space_before=Pt(0)
# Clear inherited theme-font and title-border residue in the Word base template.
for st in styles:
 for e in list(st.element.iter()):
  if e.tag==qn('w:pBdr'):e.getparent().remove(e)
  if e.tag==qn('w:rFonts'):
   for key in list(e.attrib):
    if 'theme' in key.lower():del e.attrib[key]
styles['Normal'].paragraph_format.widow_control=True
# A quiet manuscript page has only a consistent numbered footer.
h=sec.header.paragraphs[0];h.text=''
f=sec.footer.paragraphs[0];f.alignment=WD_ALIGN_PARAGRAPH.RIGHT
r=f.add_run('Master manuscript  |  ');r.font.size=Pt(9)
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');f._p.append(field)
doc.core_properties.title='Conserved genomic content within an expansive Pseudomonas aeruginosa pangenome'
doc.core_properties.subject='Evidence-constrained reconstruction of the preserved Pseudomonas genomics project'
doc.core_properties.author='Authors listed in manuscript; confirmation required'

def hyperlink(p,label,url):
 h=OxmlElement('w:hyperlink');h.set(qn('r:id'),p.part.relate_to(url,RT.HYPERLINK,is_external=True));r=OxmlElement('w:r');pr=OxmlElement('w:rPr');font=OxmlElement('w:rFonts');font.set(qn('w:ascii'),'Times New Roman');font.set(qn('w:hAnsi'),'Times New Roman');pr.append(font);col=OxmlElement('w:color');col.set(qn('w:val'),'1E4566');pr.append(col);r.append(pr);t=OxmlElement('w:t');t.text=label;r.append(t);h.append(r);p._p.append(h)
def inline(p,text):
 tokens=re.split(r'(\[[^\]]+\]\([^\)]+\)|\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)',text)
 for s in tokens:
  if not s:continue
  m=re.fullmatch(r'\[([^\]]+)\]\(([^\)]+)\)',s)
  if m:hyperlink(p,m.group(1),m.group(2));continue
  if not s.startswith(('*','`')):
   # Preserve scientific powers/subscripts as native Word run formatting.
   bits=re.split(r'(\^\(γ−1\)|\^γ|\b(?:f_g|M_i|M_j|S_ij|Σ_g)\b)',s)
   for bit in bits:
    if not bit:continue
    if bit.startswith('^'):
     r=p.add_run('γ−1' if '(' in bit else 'γ');r.font.superscript=True
    elif re.fullmatch(r'(?:f_g|M_i|M_j|S_ij|Σ_g)',bit):
     base,low=bit.split('_');p.add_run(base);p.add_run(low).font.subscript=True
    else:p.add_run(bit)
   continue
  r=p.add_run(s[2:-2] if s.startswith('**') else s[1:-1])
  if s.startswith('**'):r.bold=True
  elif s.startswith('*'):r.italic=True
  elif s.startswith('`'):r.font.name='DejaVu Sans Mono';r.font.size=Pt(9)

def m(tag,*children,**attrs):
 e=OxmlElement('m:'+tag)
 for k,v in attrs.items():e.set(qn('m:'+k),str(v))
 for c in children:e.append(c)
 return e
def mr(t):return m('r',m('t',**{})) if False else mathrun(t)
def mathrun(t):
 e=m('r');pr=m('rPr',m('sty',val='p'));e.append(pr);z=m('t');z.text=t;e.append(z);return e
def seq(*xs):return list(xs)
def sub(base,low):return m('sSub',m('e',mr(base)),m('sub',mr(low)))
def fraction(num,den,nobar=False):
 e=m('f');
 if nobar:e.append(m('fPr',m('type',val='noBar')))
 e.append(m('num',*num));e.append(m('den',*den));return e
def paren(*xs):return m('d',m('dPr',m('begChr',val='('),m('endChr',val=')')),m('e',*xs))
def binom(num,den):return paren(fraction(num,den,True))
def summation(lower,expr):return m('nary',m('naryPr',m('chr',val='∑'),m('limLoc',val='subSup'),m('supHide',val='1')),m('sub',mr(lower)),m('sup'),m('e',*expr))
def equation(which):
 p=doc.add_paragraph();p.paragraph_format.space_before=Pt(7);p.paragraph_format.space_after=Pt(9);om=m('oMath')
 if which==1:
  om.append(mr('E[P(k)] = '));b1=binom([mr('N − '),sub('f','g')],[mr('k')]);b2=binom([mr('N')],[mr('k')]);om.append(summation('g',[mr('['),mr('1 − '),fraction([b1],[b2]),mr(']')]))
 else:
  om.append(sub('φ','ij'));om.append(mr(' = '));om.append(fraction([sub('M','i'),mr(' + '),sub('M','j'),mr(' − 2'),sub('S','ij')],[sub('M','i'),mr(' + '),sub('M','j')]))
  om.append(mr(' ,    '));om.append(m('acc',m('accPr',m('chr',val='̅')),m('e',mr('φ'))));om.append(mr(' = '));om.append(fraction([mr('2')],[mr('N'),paren(mr('N − 1'))]));om.append(summation('i < j',[sub('φ','ij')]))
 p._p.append(m('oMathPara',om))

def table(lines):
 parsed=[[c.strip() for c in ln.strip().strip('|').split('|')] for ln in lines if not re.match(r'^\|\s*[-:]+',ln)]
 t=doc.add_table(rows=1,cols=len(parsed[0]));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
 widths=[1.05,.85,.75,.86,.93,.83,.85,.78] if len(parsed[0])==8 else [1.15,1.35,.95,3.45]
 ratio=6.9/sum(widths);widths=[x*ratio for x in widths]
 for j,c in enumerate(t.columns):c.width=Inches(widths[j])
 for i,row in enumerate(parsed):
  cells=t.rows[0].cells if i==0 else t.add_row().cells
  for j,(c,value) in enumerate(zip(cells,row)):
   c.width=Inches(widths[j]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER;tcpr=c._tc.get_or_add_tcPr()
   borders=OxmlElement('w:tcBorders')
   for side in ['top','left','bottom','right']:
    b=OxmlElement('w:'+side);b.set(qn('w:val'),'single');b.set(qn('w:sz'),'4');b.set(qn('w:color'),'D9D9D9');borders.append(b)
   tcpr.append(borders);margin=OxmlElement('w:tcMar')
   for side,val in [('top','80'),('bottom','80'),('left','70'),('right','70')]:
    b=OxmlElement('w:'+side);b.set(qn('w:w'),val);b.set(qn('w:type'),'dxa');margin.append(b)
   tcpr.append(margin);shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E9EDF0' if i==0 else ('F7F8F9' if i%2==0 else 'FFFFFF'));tcpr.append(shade)
   p=c.paragraphs[0];p.paragraph_format.line_spacing=1.0;p.paragraph_format.space_after=Pt(0);p.paragraph_format.keep_with_next=False;p.alignment=WD_ALIGN_PARAGRAPH.LEFT if j==0 or len(cells)==4 and j==3 else WD_ALIGN_PARAGRAPH.CENTER
   inline(p,value)
   for r in p.runs:r.font.size=Pt(9.3 if len(cells)==8 else 10);r.bold=(i==0)
  trpr=t.rows[i]._tr.get_or_add_trPr();cant=OxmlElement('w:cantSplit');trpr.append(cant)
  if i==0:repeat=OxmlElement('w:tblHeader');trpr.append(repeat)
 doc.add_paragraph().paragraph_format.space_after=Pt(1)

lines=(O/'13_MANUSCRIPT_MASTER.md').read_text().splitlines();i=0;eq=0
while i<len(lines):
 line=lines[i]
 if not line.strip():i+=1;continue
 if line.startswith('$$'):
  eq+=1;equation(eq);i+=1;continue
 if line.startswith('|'):
  block=[]
  while i<len(lines) and lines[i].startswith('|'):block.append(lines[i]);i+=1
  table(block);continue
 if line.startswith('# '):p=doc.add_paragraph(style='Title');inline(p,line[2:]);i+=1;continue
 if line.startswith('## '):
  title=line[3:]
  if title in ['Introduction','Results','Discussion','Methods','Main tables','Figure legends']:doc.add_page_break()
  p=doc.add_paragraph(style='Heading 1');inline(p,title);i+=1;continue
 if line.startswith('### '):p=doc.add_paragraph(style='Heading 2');inline(p,line[4:]);i+=1;continue
 p=doc.add_paragraph();inline(p,line)
 if re.match(r'^\d+\. ',line) and '## References' in '\n'.join(lines[:i]):p.paragraph_format.left_indent=Inches(.18);p.paragraph_format.first_line_indent=Inches(-.18)
 i+=1
assert eq==2,eq
dest=O/'14_MANUSCRIPT_CLEAN.docx';doc.save(dest)
(O/'evidence_extracts/docx_build_metadata.json').write_text(json.dumps({'input':'13_MANUSCRIPT_MASTER.md','output':dest.name,'native_display_equations':eq,'figure_embedding':'Separate vector figures; manuscript includes captions','paragraphs':len(doc.paragraphs),'tables':len(doc.tables),'page_size':'Letter portrait','authoring':'bundled python-docx; native OMML'},indent=2))
print(dest)
